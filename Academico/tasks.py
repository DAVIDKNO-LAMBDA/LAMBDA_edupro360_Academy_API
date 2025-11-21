"""
HU-11: Recordatorios Automáticos
HU-12: Consolidado Mensual de Notas

Tareas Celery para automatización del sistema académico EduPro 360.
"""
import logging
from datetime import datetime, timedelta
from django.utils import timezone
from django.core.mail import send_mail
from django.template.loader import render_to_string
from django.conf import settings
from django.contrib.auth import get_user_model
from django.db.models import Avg, Count, Q
from celery import shared_task
import os
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter, A4
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from io import BytesIO
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill

logger = logging.getLogger(__name__)
User = get_user_model()

# =============================================================================
# HU-11: RECORDATORIOS AUTOMÁTICOS
# =============================================================================

@shared_task(bind=True, max_retries=3)
def enviar_recordatorio_tareas_3_dias(self):
    """
    HU-11.1a: Envía recordatorios automáticos de tareas próximas a vencer (3 días antes).
    Se ejecuta diariamente a las 8:00 AM.
    """
    try:
        from .models import Tarea, InscripcionAsignatura
        
        # Calcular fecha límite (72 horas desde ahora)
        limite_72h_inicio = timezone.now() + timedelta(hours=72)
        limite_72h_fin = timezone.now() + timedelta(hours=96)  # 4 días para rango
        
        # Buscar tareas que vencen en aproximadamente 3 días
        tareas_3_dias = Tarea.objects.filter(
            fecha_vencimiento__gte=limite_72h_inicio,
            fecha_vencimiento__lte=limite_72h_fin,
            estado=True
        ).select_related('asignatura', 'asignatura__docente_responsable')
        
        enviados = 0
        errores = 0
        
        for tarea in tareas_3_dias:
            # Obtener estudiantes inscritos en la asignatura
            estudiantes = User.objects.filter(
                inscripciones_asignaturas__asignatura=tarea.asignatura,
                inscripciones_asignaturas__estado=True,
                groups__name='Estudiantes',
                is_active=True
            ).distinct()
            
            for estudiante in estudiantes:
                # Verificar si ya entregó la tarea
                from .models import EntregaTarea
                entrega_existente = EntregaTarea.objects.filter(
                    tarea=tarea,
                    estudiante=estudiante,
                    estado=True
                ).exists()
                
                if not entrega_existente:
                    # Enviar recordatorio temprano
                    try:
                        dias_restantes = (tarea.fecha_vencimiento - timezone.now()).days
                        tiempo_restante = tarea.fecha_vencimiento - timezone.now()
                        horas_restantes = int(tiempo_restante.total_seconds() / 3600)
                        contexto = {
                            'estudiante': estudiante,
                            'tarea': tarea,
                            'asignatura': tarea.asignatura,
                            'tiempo_restante': tiempo_restante,
                            'docente': tarea.asignatura.docente_responsable,
                            'dias_restantes': dias_restantes,
                            'horas_restantes': horas_restantes,
                            'es_recordatorio_temprano': True
                        }
                        
                        asunto = f"📅 Recordatorio Temprano: '{tarea.titulo}' vence en {dias_restantes} días"
                        mensaje_html = render_to_string('academico/recordatorio_tarea.html', contexto)
                        
                        send_mail(
                            subject=asunto,
                            message='',
                            from_email=settings.DEFAULT_FROM_EMAIL,
                            recipient_list=[estudiante.correo],
                            html_message=mensaje_html,
                            fail_silently=False
                        )
                        
                        enviados += 1
                        logger.info(f"Recordatorio 3 días enviado a {estudiante.correo} para tarea {tarea.titulo}")
                        
                    except Exception as e:
                        errores += 1
                        logger.error(f"Error enviando recordatorio 3 días: {str(e)}")
        
        resultado = f"Recordatorios 3 días: {enviados} enviados, {errores} errores"
        logger.info(resultado)
        return resultado
        
    except Exception as e:
        logger.error(f"Error en enviar_recordatorio_tareas_3_dias: {str(e)}")
        raise self.retry(countdown=300, exc=e)

@shared_task(bind=True, max_retries=3)
def enviar_recordatorio_tareas_proximas(self):
    """
    HU-11.1b: Envía recordatorios automáticos de tareas próximas a vencer (24h antes).
    Se ejecuta diariamente a las 9:00 AM.
    """
    try:
        from .models import Tarea, InscripcionAsignatura
        
        # Calcular fecha límite (24 horas desde ahora)
        limite_24h = timezone.now() + timedelta(hours=24)
        
        # Buscar tareas que vencen en las próximas 24 horas
        tareas_proximas = Tarea.objects.filter(
            fecha_vencimiento__lte=limite_24h,
            fecha_vencimiento__gt=timezone.now(),
            estado=True
        ).select_related('asignatura', 'asignatura__docente_responsable')
        
        enviados = 0
        errores = 0
        
        for tarea in tareas_proximas:
            # Obtener estudiantes inscritos en la asignatura
            estudiantes = User.objects.filter(
                inscripciones_asignaturas__asignatura=tarea.asignatura,
                inscripciones_asignaturas__estado=True,
                groups__name='Estudiantes',
                is_active=True
            ).distinct()
            
            for estudiante in estudiantes:
                # Verificar si ya entregó la tarea
                from .models import EntregaTarea
                entrega_existente = EntregaTarea.objects.filter(
                    tarea=tarea,
                    estudiante=estudiante,
                    estado=True
                ).exists()
                
                if not entrega_existente:
                    # Enviar recordatorio
                    try:
                        tiempo_restante = tarea.fecha_vencimiento - timezone.now()
                        horas_restantes = int(tiempo_restante.total_seconds() / 3600)
                        contexto = {
                            'estudiante': estudiante,
                            'tarea': tarea,
                            'asignatura': tarea.asignatura,
                            'tiempo_restante': tiempo_restante,
                            'horas_restantes': horas_restantes,
                            'docente': tarea.asignatura.docente_responsable
                        }
                        
                        asunto = f"🔔 Recordatorio: Tarea '{tarea.titulo}' vence pronto"
                        mensaje_html = render_to_string('academico/recordatorio_tarea.html', contexto)
                        
                        send_mail(
                            subject=asunto,
                            message='',  # Mensaje de texto plano vacío
                            from_email=settings.DEFAULT_FROM_EMAIL,
                            recipient_list=[estudiante.correo],
                            html_message=mensaje_html,
                            fail_silently=False
                        )
                        
                        enviados += 1
                        logger.info(f"Recordatorio enviado a {estudiante.correo} para tarea {tarea.titulo}")
                        
                    except Exception as e:
                        errores += 1
                        logger.error(f"Error enviando recordatorio a {estudiante.correo}: {str(e)}")
        
        mensaje = f"Recordatorios procesados: {enviados} enviados, {errores} errores"
        logger.info(mensaje)
        return mensaje
        
    except Exception as exc:
        logger.error(f"Error en enviar_recordatorio_tareas_proximas: {str(exc)}")
        raise self.retry(exc=exc, countdown=300, max_retries=3)


@shared_task(bind=True, max_retries=3)
def enviar_recordatorio_tareas_vencidas(self):
    """
    HU-11.2: Envía recordatorios de tareas vencidas no entregadas.
    Se ejecuta diariamente a las 10:00 AM.
    """
    try:
        from .models import Tarea, InscripcionAsignatura, EntregaTarea
        
        # Buscar tareas vencidas (hasta 7 días atrás para evitar spam)
        hace_7_dias = timezone.now() - timedelta(days=7)
        
        tareas_vencidas = Tarea.objects.filter(
            fecha_vencimiento__lt=timezone.now(),
            fecha_vencimiento__gte=hace_7_dias,
            estado=True
        ).select_related('asignatura', 'asignatura__docente_responsable')
        
        enviados = 0
        errores = 0
        
        for tarea in tareas_vencidas:
            # Obtener estudiantes que NO han entregado la tarea
            estudiantes_sin_entrega = User.objects.filter(
                inscripciones_estudiante__asignatura=tarea.asignatura,
                inscripciones_estudiante__estado=True,
                groups__name='Estudiantes',
                is_active=True
            ).exclude(
                entregas_estudiante__tarea=tarea,
                entregas_estudiante__estado=True
            ).distinct()
            
            for estudiante in estudiantes_sin_entrega:
                try:
                    dias_vencida = (timezone.now() - tarea.fecha_vencimiento).days
                    
                    contexto = {
                        'estudiante': estudiante,
                        'tarea': tarea,
                        'asignatura': tarea.asignatura,
                        'dias_vencida': dias_vencida,
                        'docente': tarea.asignatura.docente_responsable
                    }
                    
                    asunto = f"⚠️ Tarea Vencida: '{tarea.titulo}' - {dias_vencida} días de retraso"
                    mensaje_html = render_to_string('academico/recordatorio_vencida.html', contexto)
                    
                    send_mail(
                        subject=asunto,
                        message='',
                        from_email=settings.DEFAULT_FROM_EMAIL,
                        recipient_list=[estudiante.correo],
                        html_message=mensaje_html,
                        fail_silently=False
                    )
                    
                    enviados += 1
                    logger.info(f"Recordatorio de vencida enviado a {estudiante.correo} para tarea {tarea.titulo}")
                    
                except Exception as e:
                    errores += 1
                    logger.error(f"Error enviando recordatorio de vencida a {estudiante.correo}: {str(e)}")
        
        mensaje = f"Recordatorios de vencidas procesados: {enviados} enviados, {errores} errores"
        logger.info(mensaje)
        return mensaje
        
    except Exception as exc:
        logger.error(f"Error en enviar_recordatorio_tareas_vencidas: {str(exc)}")
        raise self.retry(exc=exc, countdown=300, max_retries=3)


@shared_task(bind=True, max_retries=3)
def recordar_calificaciones_pendientes(self):
    """
    HU-11.3: Envía recordatorios a docentes sobre calificaciones pendientes.
    Se ejecuta diariamente a las 8:00 AM.
    """
    try:
        from .models import EntregaTarea, CalificacionTarea
        
        # Buscar entregas sin calificar (más de 3 días de antigüedad)
        hace_3_dias = timezone.now() - timedelta(days=3)
        
        entregas_pendientes = EntregaTarea.objects.filter(
            creado__lt=hace_3_dias,
            estado=True
        ).exclude(
            id__in=CalificacionTarea.objects.filter(
                estado=True
            ).values_list('tarea_id', flat=True)
        ).select_related('tarea', 'tarea__asignatura', 'tarea__asignatura__docente_responsable', 'estudiante')
        
        # Agrupar por docente
        docentes_entregas = {}
        for entrega in entregas_pendientes:
            docente = entrega.tarea.asignatura.docente_responsable
            if docente not in docentes_entregas:
                docentes_entregas[docente] = []
            docentes_entregas[docente].append(entrega)
        
        enviados = 0
        errores = 0
        
        for docente, entregas in docentes_entregas.items():
            try:
                contexto = {
                    'docente': docente,
                    'entregas_pendientes': entregas,
                    'total_pendientes': len(entregas),
                    'asignaturas': list(set([e.tarea.asignatura for e in entregas]))
                }
                
                asunto = f"📝 Recordatorio: {len(entregas)} calificaciones pendientes"
                mensaje_html = render_to_string('academico/recordatorio_calificaciones.html', contexto)
                
                send_mail(
                    subject=asunto,
                    message='',
                    from_email=settings.DEFAULT_FROM_EMAIL,
                    recipient_list=[docente.correo],
                    html_message=mensaje_html,
                    fail_silently=False
                )
                
                enviados += 1
                logger.info(f"Recordatorio de calificaciones enviado a {docente.correo}")
                
            except Exception as e:
                errores += 1
                logger.error(f"Error enviando recordatorio de calificaciones a {docente.correo}: {str(e)}")
        
        mensaje = f"Recordatorios de calificaciones procesados: {enviados} enviados, {errores} errores"
        logger.info(mensaje)
        return mensaje
        
    except Exception as exc:
        logger.error(f"Error en recordar_calificaciones_pendientes: {str(exc)}")
        raise self.retry(exc=exc, countdown=300, max_retries=3)


# =============================================================================
# HU-12: CONSOLIDADO MENSUAL DE NOTAS
# =============================================================================

@shared_task(bind=True, max_retries=3)
def generar_consolidado_mensual(self, mes=None, año=None):
    """
    HU-12.1: Genera consolidado mensual de notas en PDF y Excel.
    Se ejecuta automáticamente el día 5 de cada mes a las 6:00 AM.
    """
    try:
        from .models import CalificacionTarea, Tarea, InscripcionAsignatura, PeriodoAcademico
        
        # Si no se especifica mes/año, usar el mes anterior
        if not mes or not año:
            fecha_anterior = timezone.now() - timedelta(days=30)
            mes = fecha_anterior.month
            año = fecha_anterior.year
        
        # Obtener calificaciones del mes
        calificaciones = CalificacionTarea.objects.filter(
            creado__year=año,
            creado__month=mes,
            estado=True
        ).select_related(
            'tarea', 'tarea__asignatura', 'tarea__asignatura__periodo_academico',
            'estudiante'
        ).order_by('tarea__asignatura__nombre', 'estudiante__apellidos')
        
        if not calificaciones.exists():
            logger.info(f"No hay calificaciones para el mes {mes}/{año}")
            return f"No hay datos para procesar del mes {mes}/{año}"
        
        # Generar PDF
        pdf_path = generar_pdf_consolidado(calificaciones, mes, año)
        
        # Generar Excel
        excel_path = generar_excel_consolidado(calificaciones, mes, año)
        
        # Enviar por correo a administradores
        enviar_consolidado_email(pdf_path, excel_path, mes, año)
        
        mensaje = f"Consolidado mensual generado para {mes}/{año}: PDF y Excel creados y enviados"
        logger.info(mensaje)
        return mensaje
        
    except Exception as exc:
        logger.error(f"Error en generar_consolidado_mensual: {str(exc)}")
        raise self.retry(exc=exc, countdown=300, max_retries=3)


def generar_pdf_consolidado(calificaciones, mes, año):
    """Genera reporte PDF del consolidado mensual."""
    try:
        # Crear directorio si no existe
        reports_dir = os.path.join(settings.MEDIA_ROOT, 'reportes')
        os.makedirs(reports_dir, exist_ok=True)
        
        filename = f"consolidado_{año}_{mes:02d}.pdf"
        filepath = os.path.join(reports_dir, filename)
        
        # Crear documento PDF
        doc = SimpleDocTemplate(filepath, pagesize=A4)
        styles = getSampleStyleSheet()
        story = []
        
        # Título
        titulo_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=16,
            spaceAfter=30,
            alignment=1  # Centro
        )
        
        meses = [
            '', 'Enero', 'Febrero', 'Marzo', 'Abril', 'Mayo', 'Junio',
            'Julio', 'Agosto', 'Septiembre', 'Octubre', 'Noviembre', 'Diciembre'
        ]
        
        titulo = f"Consolidado Mensual de Notas<br/>{meses[mes]} {año}"
        story.append(Paragraph(titulo, titulo_style))
        story.append(Spacer(1, 20))
        
        # Estadísticas generales
        total_calificaciones = calificaciones.count()
        promedio_general = calificaciones.aggregate(Avg('nota'))['nota__avg'] or 0
        
        stats_text = f"""
        <b>Estadísticas Generales:</b><br/>
        • Total de calificaciones: {total_calificaciones}<br/>
        • Promedio general: {promedio_general:.2f}<br/>
        • Período: {meses[mes]} {año}
        """
        story.append(Paragraph(stats_text, styles['Normal']))
        story.append(Spacer(1, 20))
        
        # Tabla de calificaciones por asignatura
        asignaturas = {}
        for cal in calificaciones:
            asig_name = cal.tarea.asignatura.nombre
            if asig_name not in asignaturas:
                asignaturas[asig_name] = []
            asignaturas[asig_name].append(cal)
        
        for asignatura, cals in asignaturas.items():
            # Título de asignatura
            asig_title = f"<b>{asignatura}</b>"
            story.append(Paragraph(asig_title, styles['Heading2']))
            
            # Datos de la tabla
            data = [['Estudiante', 'Tarea', 'Nota', 'Fecha']]
            
            for cal in cals:
                data.append([
                    f"{cal.estudiante.nombres} {cal.estudiante.apellidos}".title(),
                    cal.tarea.titulo,
                    f"{cal.nota:.1f}",
                    cal.creado.strftime('%d/%m/%Y')
                ])
            
            # Crear tabla
            table = Table(data)
            table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 10),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
                ('GRID', (0, 0), (-1, -1), 1, colors.black)
            ]))
            
            story.append(table)
            story.append(Spacer(1, 20))
        
        # Generar PDF
        doc.build(story)
        logger.info(f"PDF generado: {filepath}")
        return filepath
        
    except Exception as e:
        logger.error(f"Error generando PDF: {str(e)}")
        raise


def generar_excel_consolidado(calificaciones, mes, año):
    """Genera reporte Excel del consolidado mensual."""
    try:
        # Crear directorio si no existe
        reports_dir = os.path.join(settings.MEDIA_ROOT, 'reportes')
        os.makedirs(reports_dir, exist_ok=True)
        
        filename = f"consolidado_{año}_{mes:02d}.xlsx"
        filepath = os.path.join(reports_dir, filename)
        
        # Crear workbook
        wb = openpyxl.Workbook()
        
        # Hoja de resumen
        ws_resumen = wb.active
        ws_resumen.title = "Resumen"
        
        # Configurar estilos
        header_font = Font(bold=True, size=12)
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        center_alignment = Alignment(horizontal="center", vertical="center")
        
        # Título
        ws_resumen['A1'] = f"Consolidado Mensual - {mes:02d}/{año}"
        ws_resumen['A1'].font = Font(bold=True, size=16)
        ws_resumen.merge_cells('A1:D1')
        
        # Estadísticas
        ws_resumen['A3'] = "Total Calificaciones:"
        ws_resumen['B3'] = calificaciones.count()
        ws_resumen['A4'] = "Promedio General:"
        ws_resumen['B4'] = round(calificaciones.aggregate(Avg('nota'))['nota__avg'] or 0, 2)
        
        # Hoja detallada
        ws_detalle = wb.create_sheet("Detalle Calificaciones")
        
        # Headers
        headers = ['Estudiante', 'Asignatura', 'Tarea', 'Tipo', 'Nota', 'Peso %', 'Fecha']
        for col, header in enumerate(headers, 1):
            cell = ws_detalle.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_alignment
        
        # Datos
        for row, cal in enumerate(calificaciones, 2):
            ws_detalle.cell(row=row, column=1, value=f"{cal.estudiante.nombres} {cal.estudiante.apellidos}".title())
            ws_detalle.cell(row=row, column=2, value=cal.tarea.asignatura.nombre)
            ws_detalle.cell(row=row, column=3, value=cal.tarea.titulo)
            ws_detalle.cell(row=row, column=4, value=cal.tarea.tipo_tarea)
            ws_detalle.cell(row=row, column=5, value=float(cal.nota))
            ws_detalle.cell(row=row, column=6, value=float(cal.tarea.peso_porcentual))
            ws_detalle.cell(row=row, column=7, value=cal.creado.strftime('%d/%m/%Y'))
        
        # Ajustar ancho de columnas
        for column in ws_detalle.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws_detalle.column_dimensions[column_letter].width = adjusted_width
        
        # Guardar archivo
        wb.save(filepath)
        logger.info(f"Excel generado: {filepath}")
        return filepath
        
    except Exception as e:
        logger.error(f"Error generando Excel: {str(e)}")
        raise


def enviar_consolidado_email(pdf_path, excel_path, mes, año):
    """Envía el consolidado por email a administradores."""
    try:
        # Obtener usuarios con permiso para recibir reportes mensuales
        administradores = User.objects.filter(
            user_permissions__codename='recibir_notificacion_estado_mensual',
            is_active=True
        ).distinct()
        
        # Si no hay usuarios con el permiso específico, usar administradores como fallback
        if not administradores.exists():
            administradores = User.objects.filter(
                groups__name='Administradores',
                is_active=True
            )
        
        if not administradores.exists():
            logger.warning("No hay administradores para enviar consolidado")
            return
        
        meses = [
            '', 'Enero', 'Febrero', 'Marzo', 'Abril', 'Mayo', 'Junio',
            'Julio', 'Agosto', 'Septiembre', 'Octubre', 'Noviembre', 'Diciembre'
        ]
        
        contexto = {
            'mes': meses[mes],
            'año': año,
            'fecha_generacion': timezone.now()
        }
        
        asunto = f"📊 Consolidado Mensual {meses[mes]} {año} - EduPro 360"
        mensaje_html = render_to_string('academico/consolidado_mensual.html', contexto)
        
        from django.core.mail import EmailMessage
        
        for admin in administradores:
            email = EmailMessage(
                subject=asunto,
                body=mensaje_html,
                from_email=settings.DEFAULT_FROM_EMAIL,
                to=[admin.correo]
            )
            email.content_subtype = "html"
            
            # Adjuntar archivos
            if os.path.exists(pdf_path):
                email.attach_file(pdf_path)
            if os.path.exists(excel_path):
                email.attach_file(excel_path)
            
            email.send()
            logger.info(f"Consolidado enviado a {admin.correo}")
        
    except Exception as e:
        logger.error(f"Error enviando consolidado por email: {str(e)}")
        raise


# =============================================================================
# TAREAS DE ESTADÍSTICAS ADICIONALES
# =============================================================================

@shared_task
def generar_estadisticas_semanales():
    """Genera estadísticas semanales para dashboard."""
    try:
        from .models import CalificacionTarea, EntregaTarea
        
        # Última semana
        hace_semana = timezone.now() - timedelta(days=7)
        
        stats = {
            'calificaciones_semana': CalificacionTarea.objects.filter(
                creado__gte=hace_semana,
                estado=True
            ).count(),
            'entregas_semana': EntregaTarea.objects.filter(
                creado__gte=hace_semana,
                estado=True
            ).count(),
            'promedio_semanal': CalificacionTarea.objects.filter(
                creado__gte=hace_semana,
                estado=True
            ).aggregate(Avg('nota'))['nota__avg'] or 0
        }
        
        logger.info(f"Estadísticas semanales: {stats}")
        return stats
        
    except Exception as e:
        logger.error(f"Error generando estadísticas semanales: {str(e)}")
        raise


@shared_task
def limpiar_archivos_antiguos():
    """Limpia archivos de reportes antiguos (más de 6 meses)."""
    try:
        reports_dir = os.path.join(settings.MEDIA_ROOT, 'reportes')
        if not os.path.exists(reports_dir):
            return "No hay directorio de reportes"
        
        hace_6_meses = timezone.now() - timedelta(days=180)
        archivos_eliminados = 0
        
        for filename in os.listdir(reports_dir):
            filepath = os.path.join(reports_dir, filename)
            if os.path.isfile(filepath):
                # Obtener fecha de modificación
                fecha_mod = datetime.fromtimestamp(os.path.getmtime(filepath))
                fecha_mod = timezone.make_aware(fecha_mod)
                
                if fecha_mod < hace_6_meses:
                    os.remove(filepath)
                    archivos_eliminados += 1
                    logger.info(f"Archivo eliminado: {filename}")
        
        mensaje = f"Limpieza completada: {archivos_eliminados} archivos eliminados"
        logger.info(mensaje)
        return mensaje
        
    except Exception as e:
        logger.error(f"Error limpiando archivos antiguos: {str(e)}")
        raise